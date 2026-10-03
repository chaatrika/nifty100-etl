# run_screener.py
# Runs the 6 preset screeners (or your own thresholds) and writes output/screener_output.xlsx
# (Sprint 3, Day 16 and 17)
#
# Run from the project root:
#   python -m src.screener.run_screener
#   python -m src.screener.run_screener --custom roe_min=18 de_max=0.5

import argparse
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import engine as E

OUT_FILE = os.path.join("output", "screener_output.xlsx")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
HEADER = PatternFill("solid", fgColor="1F3864")

# which KPI column each filter is checked on (for the colours)
SPECIAL_COLUMN = {
    "fcf_positive_latest": "free_cash_flow_cr",
    "de_declining": "debt_to_equity",
}


def filter_column(name):
    if name in SPECIAL_COLUMN:
        return SPECIAL_COLUMN[name]
    return E.FILTERS[name][0]


def nice_value(row, column):
    # what to write in the cell
    if column == "interest_coverage" and row.get("icr_label") == "Debt Free":
        return "Debt Free"
    value = row[column]
    if value is None or value != value:  # empty / NaN
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    return value


def write_sheet(wb, title, filters, data, checks, n_pass):
    ws = wb.create_sheet(title[:31])
    columns = ["company_id", "company_name", "broad_sector"] + E.KPI_COLUMNS

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    filter_text = ", ".join(str(k) + " = " + str(v) for k, v in filters.items())
    ws.cell(
        row=2,
        column=1,
        value="Filters: "
        + filter_text
        + "   (green = meets the filter, red = fails it)",
    )

    header_row = 4
    for c, name in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    row_number = header_row + 1
    for i in range(len(data)):
        # a blank line and a small title before the "near miss" companies
        if i == n_pass and len(data) > n_pass:
            row_number += 1
            ws.cell(
                row=row_number,
                column=1,
                value="Near misses - companies that fail exactly one filter",
            ).font = Font(bold=True, italic=True)
            row_number += 1

        row = data.iloc[i].to_dict()
        for c, name in enumerate(columns, start=1):
            ws.cell(row=row_number, column=c, value=nice_value(row, name))

        # colour the cells that belong to a filter
        for filter_name in filters:
            column = filter_column(filter_name)
            if column not in columns:
                continue
            passed = bool(checks.iloc[i][filter_name])
            ws.cell(row=row_number, column=columns.index(column) + 1).fill = (
                GREEN if passed else RED
            )
        row_number += 1

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.row_dimensions[header_row].height = 45
    for c, name in enumerate(columns, start=1):
        width = 30 if name == "company_name" else 16
        ws.column_dimensions[get_column_letter(c)].width = width


def sheet_data(df, filters, exclude_sectors, near_miss_count=10):
    # passing companies first, then a few near misses (they fail exactly one filter)
    data = df
    if exclude_sectors:
        data = data[~data["broad_sector"].isin(exclude_sectors)]
    checks = E.check_filters(data, filters)
    passed = checks.all(axis=1)
    failed_count = (~checks).sum(axis=1)

    pass_rows = data[passed].sort_values("composite_quality_score", ascending=False)
    near = (
        data[(failed_count == 1)]
        .sort_values("composite_quality_score", ascending=False)
        .head(near_miss_count)
    )

    combined = pd.concat([pass_rows, near])
    return combined, checks.loc[combined.index], len(pass_rows)


def write_excel(df, config, path=OUT_FILE):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}
    for name, preset in config["presets"].items():
        data, checks, n_pass = sheet_data(
            df, preset["filters"], preset.get("exclude_sectors")
        )
        write_sheet(wb, preset["title"], preset["filters"], data, checks, n_pass)
        counts[name] = n_pass
    wb.save(path)
    return counts


def parse_custom(items):
    # turns ["roe_min=18", "de_max=0.5"] into {"roe_min": 18.0, "de_max": 0.5}
    filters = {}
    for item in items:
        key, value = item.split("=")
        if value.lower() in ("true", "false"):
            filters[key] = value.lower() == "true"
        else:
            filters[key] = float(value)
    return filters


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--custom", nargs="+", help="your own thresholds, e.g. roe_min=18 de_max=0.5"
    )
    args = parser.parse_args()

    df = E.get_universe()
    config = E.load_config()

    if args.custom:
        result = E.run_custom(df, parse_custom(args.custom))
        print("Custom screen returned", len(result), "companies")
        print(
            result[["company_id", "company_name", "composite_quality_score"]].to_string(
                index=False
            )
        )
        return

    counts = write_excel(df, config)
    print("Preset results (need 5 to 50 each):")
    for name, n in counts.items():
        status = "OK" if 5 <= n <= 50 else "CHECK"
        print("  %-22s %3d companies  %s" % (name, n, status))
    print("Written:", OUT_FILE)


if __name__ == "__main__":
    main()
