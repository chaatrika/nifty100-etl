# valuation.py
# FCF yield + P/E based overvaluation flags for every company (Sprint 4, Day 26)
#
# Run from the project root:   python -m src.analytics.valuation
#
# Reads  : data/nifty100.db  (market_cap, financial_ratios, companies, sectors)
# Writes : output/valuation_summary.xlsx   (all companies)
#          output/valuation_flags.csv      (only Caution / Discount companies)

import os
import sqlite3

import numpy as np
import pandas as pd

DB_PATH = os.path.join("data", "nifty100.db")
OUT_DIR = "output"
CAUTION_MULT = 1.5      # P/E above sector median x 1.5  -> Caution
DISCOUNT_MULT = 0.7     # P/E below sector median x 0.7  -> Discount

COLUMNS = ["company_id", "company_name", "sector", "P/E", "P/B", "EV/EBITDA",
           "FCF_yield_pct", "5yr_median_PE", "PE_vs_sector_median_pct", "flag"]


def flag_for(pe, sector_median):
    # Caution / Discount / Fair. If the P/E or the sector median is missing we say N/A
    # (we don't guess a label without data).
    if pd.isna(pe) or pd.isna(sector_median) or sector_median <= 0 or pe <= 0:
        return "N/A"
    if pe > sector_median * CAUTION_MULT:
        return "Caution"
    if pe < sector_median * DISCOUNT_MULT:
        return "Discount"
    return "Fair"


def build_valuation(db_path=DB_PATH):
    conn = sqlite3.connect(db_path)
    mc = pd.read_sql("SELECT * FROM market_cap", conn)
    fr = pd.read_sql("SELECT company_id, year, free_cash_flow_cr, net_profit_margin_pct "
                     "FROM financial_ratios", conn)
    names = pd.read_sql("SELECT id AS company_id, company_name FROM companies", conn)
    sectors = pd.read_sql("SELECT company_id, broad_sector AS sector FROM sectors", conn)
    conn.close()

    latest_year = int(mc["year"].max())
    latest = mc[mc["year"] == latest_year].copy()

    # FCF of the same financial year as the market cap (last row of that year that has profit data)
    fr = fr[fr["net_profit_margin_pct"].notna()].copy()
    fr["fy"] = fr["year"].str[:4].astype(int)
    fr = fr[fr["fy"] == latest_year].sort_values("year").groupby("company_id").tail(1)
    latest = latest.merge(fr[["company_id", "free_cash_flow_cr"]], on="company_id", how="left")

    # FCF yield = FCF / market cap x 100 (both in Crore)
    latest["FCF_yield_pct"] = latest["free_cash_flow_cr"] / latest["market_cap_crore"] * 100

    # 5 year median P/E (latest year and the 4 before it)
    last5 = mc[mc["year"] > latest_year - 5]
    med5 = last5.groupby("company_id")["pe_ratio"].median().rename("5yr_median_PE")
    latest = latest.merge(med5, on="company_id", how="left")

    df = latest.merge(names, on="company_id", how="left").merge(sectors, on="company_id", how="left")

    # sector median P/E in the latest year
    df["sector_median_pe"] = df.groupby("sector")["pe_ratio"].transform("median")
    df["PE_vs_sector_median_pct"] = (df["pe_ratio"] / df["sector_median_pe"] - 1) * 100
    df["flag"] = [flag_for(p, m) for p, m in zip(df["pe_ratio"], df["sector_median_pe"])]

    df = df.rename(columns={"pe_ratio": "P/E", "pb_ratio": "P/B", "ev_ebitda": "EV/EBITDA"})
    out = df[COLUMNS].copy()
    num = ["P/E", "P/B", "EV/EBITDA", "FCF_yield_pct", "5yr_median_PE", "PE_vs_sector_median_pct"]
    out[num] = out[num].astype(float).round(2)
    return out.sort_values("company_id").reset_index(drop=True)


def write_outputs(df, out_dir=OUT_DIR):
    os.makedirs(out_dir, exist_ok=True)
    xlsx = os.path.join(out_dir, "valuation_summary.xlsx")
    df.to_excel(xlsx, index=False, sheet_name="valuation")
    # small touch: widen columns and freeze the header
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(xlsx)
    ws = wb.active
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, len(str(col[0].value)) + 4)
    ws.column_dimensions["B"].width = 36
    ws.freeze_panes = "A2"
    wb.save(xlsx)

    flagged = df[df["flag"].isin(["Caution", "Discount"])].sort_values(["flag", "PE_vs_sector_median_pct"])
    flagged.to_csv(os.path.join(out_dir, "valuation_flags.csv"), index=False)
    return xlsx, flagged


if __name__ == "__main__":
    data = build_valuation()
    path, flagged = write_outputs(data)
    print("Companies:", len(data))
    print(data["flag"].value_counts().to_string())
    print("Written:", path, "and", os.path.join(OUT_DIR, "valuation_flags.csv"), f"({len(flagged)} rows)")
