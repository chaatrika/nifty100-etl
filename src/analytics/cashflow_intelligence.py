# cashflow_intelligence.py - cash flow intelligence for every company (Sprint 5, Day 31)
#
# Run from the project root:   python -m src.analytics.cashflow_intelligence
#
# Builds on the formulas already in cashflow_kpis.py (Sprint 2), which we do not change.
# Writes : output/cashflow_intelligence.xlsx   one row per company
#          output/distress_alerts.csv          CFO < 0 and CFF > 0 in the latest year

import os

import numpy as np
import pandas as pd

from src.analytics import cagr as C
from src.analytics import cashflow_kpis as K
from src.analytics import series

OUT_DIR = "output"
CAPITAL_CSV = os.path.join(OUT_DIR, "capital_allocation.csv")
COLUMNS = ["company_id", "sector", "cfo_quality_score", "cfo_quality_label", "capex_intensity_pct", "capex_label",
           "fcf_cagr_5yr", "fcf_conversion_pct", "distress_flag", "deleveraging_flag", "capital_allocation_label"]


def latest_patterns(path=CAPITAL_CSV):
    # latest annual pattern for each company
    if not os.path.exists(path):
        return {}
    df = pd.read_csv(path)
    df = series.annual(df, "pattern_label").sort_values("fy").groupby("company_id").tail(1)
    return dict(zip(df["company_id"], df["pattern_label"]))


def build(data=None):
    data = data or series.load_all()
    cf, pl, bs = (series.by_company(data[k]) for k in ("cf", "pl", "bs"))
    patterns = latest_patterns()
    empty = pd.DataFrame(columns=["fy"])
    rows, alerts = [], []
    for c in data["companies"].itertuples():
        cid = c.company_id
        cfd, pld, bsd = cf.get(cid, empty), pl.get(cid, empty), bs.get(cid, empty)
        row = dict.fromkeys(COLUMNS)
        row.update(company_id=cid, sector=c.sector, capital_allocation_label=patterns.get(cid, "N/A"),
                   distress_flag=False, deleveraging_flag=False)
        if len(cfd):
            m = cfd.merge(pld[["fy", "sales", "net_profit", "operating_profit"]], on="fy", how="left")
            m["fcf"] = m["operating_activity"] + m["investing_activity"]      # FCF = CFO + CFI
            last = m.iloc[-1]

            # CFO quality: average of CFO / PAT over the last 5 years
            score = K.cfo_quality_score(m["operating_activity"].tolist(), m["net_profit"].tolist(), 5)
            row["cfo_quality_score"] = None if score is None else round(score, 2)
            row["cfo_quality_label"] = K.cfo_quality_label(score)

            # CapEx intensity: abs(investing) / sales x 100 (latest year)
            ci = K.capex_intensity(last["investing_activity"], last["sales"])
            row["capex_intensity_pct"] = None if ci is None else round(ci, 2)
            row["capex_label"] = K.capex_label(ci)

            # 5-year FCF CAGR and FCF conversion (FCF / operating profit)
            fcf_series = {int(fy): v for fy, v in zip(m["fy"], m["fcf"]) if pd.notna(v)}
            value, _flag = C.cagr_from_series(fcf_series, 5, end_year=int(last["fy"]))
            row["fcf_cagr_5yr"] = None if value is None or pd.isna(value) else round(value, 2)
            fc = K.fcf_conversion(last["fcf"], last["operating_profit"])
            row["fcf_conversion_pct"] = None if fc is None else round(fc, 2)

            # distress: burning cash in operations while raising money from financing
            cfo, cff = K.to_num(last["operating_activity"]), K.to_num(last["financing_activity"])
            row["distress_flag"] = bool(cfo is not None and cff is not None and cfo < 0 and cff > 0)
            if row["distress_flag"]:
                alerts.append((cid, c.company_name, c.sector, int(last["fy"]), cfo, cff, K.to_num(last["net_profit"])))

            # deleveraging: paying down debt (CFF < 0) and borrowings lower than the year before
            b = bsd.set_index("fy")["borrowings"] if len(bsd) else pd.Series(dtype=float)
            fy = int(last["fy"])
            if cff is not None and cff < 0 and fy in b.index and (fy - 1) in b.index \
                    and pd.notna(b[fy]) and pd.notna(b[fy - 1]):
                row["deleveraging_flag"] = bool(b[fy] < b[fy - 1])
        rows.append(row)

    out = pd.DataFrame(rows, columns=COLUMNS)
    alert_df = pd.DataFrame(alerts, columns=["company_id", "company_name", "sector", "year", "cfo_cr", "cff_cr",
                                             "latest_net_profit_cr"])
    return out, alert_df


def write_outputs(out, alerts, out_dir=OUT_DIR):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "cashflow_intelligence.xlsx")
    out.to_excel(path, index=False, sheet_name="cashflow_intelligence")
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    wb = load_workbook(path)
    ws = wb.active
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 4)
    ws.freeze_panes = "A2"
    wb.save(path)
    alerts.to_csv(os.path.join(out_dir, "distress_alerts.csv"), index=False)
    return path


if __name__ == "__main__":
    result, alert_rows = build()
    print("Written:", write_outputs(result, alert_rows))
    print("Companies:", len(result))
    for col in ("cfo_quality_label", "capex_label", "capital_allocation_label"):
        print(result[col].value_counts(dropna=False).to_string(), "\n")
    print("Distress alerts:", len(alert_rows), "| Deleveraging:", int(result["deleveraging_flag"].sum()))
